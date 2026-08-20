"""Excel application tracker.

Appends one row per application with the columns you asked for, plus a few that
make the sheet genuinely useful later (match score, job URL, status).
"""
from __future__ import annotations
from datetime import datetime

from openpyxl import Workbook, load_workbook

from . import config

HEADERS = [
    "Company", "Job Title", "Date", "Time",
    "Resume Used", "Match Score", "Job URL", "Status",
]


def _ensure_workbook():
    path = config.TRACKER_PATH
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append(HEADERS)
    wb.save(path)
    return wb


def log_application(*, company, job_title, resume_used, match_score, job_url, status="Prepared"):
    wb = _ensure_workbook()
    ws = wb["Applications"]
    now = datetime.now()
    ws.append([
        company, job_title,
        now.strftime("%Y-%m-%d"), now.strftime("%H:%M:%S"),
        resume_used, match_score, job_url, status,
    ])
    wb.save(config.TRACKER_PATH)
    return str(config.TRACKER_PATH)
