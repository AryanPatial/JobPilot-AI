"""Manual checks. Run: python check.py free"""
import json
import sys

CHECKS = {}


def check(name, cost="FREE", note=""):
    def wrap(fn):
        CHECKS[name] = (fn, cost, note or (fn.__doc__ or "").strip().split("\n")[0])
        return fn
    return wrap


# --------------------------------------------------------------- FREE ----- #
@check("location", note="US-only filter, including the tricky cases")
def _location():
    import helpers as H
    cases = [("San Francisco, CA", True), ("Manchester, NH", True),
             ("Paris, TX", True), ("Remote - US", True),
             ("Tokyo, Japan", False), ("Bengaluru, IN", False),
             ("Berlin, DE", False), ("Toronto, CA", False),
             ("Remote", False), ("", False),
             ("New York, NY;Toronto, Ontario, CAN", True),
             ("Tokyo, Japan; London, UK", False)]
    bad = 0
    for loc, want in cases:
        got = H.is_us_location(loc)
        bad += got != want
        mark = "ok  " if got == want else "FAIL"
        shown = ", ".join(H.us_locations(loc)) or "-"
        print(f"  {mark} {loc!r:38} -> {str(got):<6} kept: {shown}")
    return bad == 0


@check("salary", note="pulls the top of a pay range out of a JD")
def _salary():
    import helpers as H
    cases = [("Range is $140,000 - $348,000 USD", "$348,000"),
             ("We pay $95K to $130K.", "$130,000"),
             ("No pay listed here.", "Negotiable")]
    bad = 0
    for jd, want in cases:
        got = H.salary_expectation_from_jd(jd)
        bad += got != want
        print(f"  {'ok  ' if got == want else 'FAIL'} {jd[:40]!r:44} -> {got}")
    return bad == 0


@check("experience", note="years of experience computed from resume dates")
def _experience():
    import helpers as H
    for v in H.load_resume_variants():
        print(f"  {v['resume_id']:<18}{H.years_of_experience(v)} years")
    return True


@check("scoring", note="deterministic JD coverage of a resume")
def _scoring():
    import helpers as H, scoring as S, requests
    vocab = H.claimable_vocabulary()
    print(f"  vocabulary: {len(vocab)} claimable terms")
    r = requests.get("https://boards-api.greenhouse.io/v1/boards/figma/jobs",
                     params={"content": "true"}, timeout=30)
    job = next(j for j in r.json()["jobs"] if j["title"] == "Data Engineer")
    jd = H._strip_html(job["content"])
    for v in H.load_resume_variants():
        sc = S.score_resume(v, jd, vocab)
        print(f"  {v['resume_id']:<18}{sc['score']:>3}/100  missing: "
              f"{', '.join(sc['missing'][:5]) or '-'}")
    return True


@check("search", note="live Greenhouse search, no LLM")
def _search():
    import helpers as H
    jobs = H.search_jobs_greenhouse(["figma", "chime"], ["Data", "Machine Learning"])
    print(f"  {len(jobs)} US jobs found")
    for j in jobs[:8]:
        print(f"    {j['title'][:46]:<48}{j['location'][:34]}")
    return True


@check("pdf", note="render a tailored resume to PDF")
def _pdf():
    import helpers as H, os
    path = H.render_resume_pdf(H.load_resume_variants()[0], "CheckScript",
                               "Machine Learning Engineer, Ranking")
    print(f"  {path}  ({os.path.getsize(path):,} bytes)")
    return os.path.getsize(path) > 10_000


@check("tracker", note="write a row to the Excel tracker")
def _tracker():
    import helpers as H
    from openpyxl import load_workbook
    H.log_application(company="CheckScript", job_title="Test Role",
                      resume_used="test.pdf", match_score=80, coverage=91,
                      attempts=2, job_url="https://example.com", status="TEST ROW")
    ws = load_workbook(H.TRACKER_PATH)["Applications"]
    for row in list(ws.iter_rows(values_only=True))[-2:]:
        print("  ", row)
    print("  (delete the TEST ROW before you use the sheet for real)")
    return True


@check("blocked", note="questions the model must never answer")
def _blocked():
    import apply
    must_block = ["Agreement to Arbitrate*", "Electronic Signature",
                  "Do you consent to a background check?", "Type your initials",
                  "I certify the above is true"]
    must_allow = ["Why Anthropic?*", "Do you have expertise coding in Python?*",
                  "How did you hear about this job?"]
    bad = 0
    for q in must_block:
        ok = bool(apply.NEVER_ANSWER.search(q))
        bad += not ok
        print(f"  {'ok  ' if ok else 'FAIL'} BLOCKED  {q}")
    for q in must_allow:
        ok = not apply.NEVER_ANSWER.search(q)
        bad += not ok
        print(f"  {'ok  ' if ok else 'FAIL'} allowed  {q}")
    return bad == 0


@check("checkpoint", note="the human pause cannot be bypassed by a pipe")
def _checkpoint():
    import run
    if sys.stdin.isatty():
        print("  running in a terminal - the checkpoint would prompt you.")
        print("  to prove it refuses a pipe:  echo | ./venv/bin/python check.py checkpoint")
        return True
    ok = run.confirm_submitted() is False
    print(f"  {'ok  ' if ok else 'FAIL'} piped stdin correctly refused to approve")
    return ok


@check("circumstances", note="facts answered from the profile, never by the model")
def _circumstances():
    """Questions with a stored answer must be resolved by rule, not guessed."""
    import llm_tasks as T
    cases = [("Have you ever interviewed at Anthropic before?*", "No"),
             ("Have you interviewed for this role or another role in the past 6 months?", "No"),
             ("Were you referred by an employee?", "No"),
             ("Have you previously applied to this company?", "No"),
             ("Do you have any other offers pending?", "No"),
             ("Are you currently employed?", "Yes")]
    bad = 0
    for q, want in cases:
        got = T._circumstance_answer(q)
        ok = got is not None and got.strip().lower().startswith(want.lower())
        bad += not ok
        print(f"  {'ok  ' if ok else 'FAIL'} {q[:56]:<58} -> {str(got)[:30]}")
    # Questions a rule must NOT claim
    for q in ["Why Anthropic?*", "Do you have expertise coding in Python?*",
              "Preferred First Name", "Preferred Name", "Referral source",
              "How did you hear about this job?*"]:
        got = T._circumstance_answer(q)
        ok = got is None
        bad += not ok
        print(f"  {'ok  ' if ok else 'FAIL'} {q[:56]:<58} -> "
              f"{'left to the model' if ok else f'WRONGLY ANSWERED {got!r}'}")
    return bad == 0


@check("fabrication", note="an invented citation must score zero")
def _fabrication():
    """The safety property of verified scoring, tested without an LLM."""
    import scoring as S
    resume = "Built a retrieval-augmented QA service over 900 internal documents"
    reqs = [
        {"requirement": "builds RAG systems", "must_have": True,
         "evidence": "Built a retrieval-augmented QA service over 900 internal documents"},
        {"requirement": "Kubernetes at scale", "must_have": True,
         "evidence": "Orchestrated 200-node Kubernetes clusters serving 4M requests"},
        {"requirement": "paraphrased, not quoted", "must_have": True,
         "evidence": "Built a RAG service over ~900 docs"},
        {"requirement": "honestly not covered", "must_have": True, "evidence": ""},
    ]
    cov = S.verify_and_score(resume, reqs)
    want = ["met", "fabricated", "fabricated", "unmet"]
    bad = 0
    for d, w in zip(cov["detail"], want):
        ok = d["verdict"] == w
        bad += not ok
        print(f"  {'ok  ' if ok else 'FAIL'} {d['requirement']:<26} -> {d['verdict']}")
    ok = cov["score"] == 25
    bad += not ok
    print(f"  {'ok  ' if ok else 'FAIL'} score {cov['score']}/100 - only the real "
          f"citation counted (expected 25)")
    return bad == 0


@check("invention", note="a term with no basis in your files is flagged")
def _invention():
    import helpers as H, scoring as S
    srcs = [H.EXPERIENCE_POOL_PATH] + H.resume_variant_paths()
    clean = H.load_resume_variants()[0]
    tampered = json.loads(json.dumps(clean))
    tampered["experience"][0]["bullets"].append(
        "Deployed Kubernetes and Terraform across a HashiCorp Vault estate")
    a = S.introduced_terms(clean, *srcs)
    b = S.introduced_terms(tampered, *srcs)
    caught = set(b) - set(a)
    print(f"  untouched résumé  -> {len(a)} flagged (want 0)")
    print(f"  after inventing   -> {sorted(caught)}")
    # kubernetes IS in the pool, so it should not be flagged
    ok = not a and {"terraform", "hashicorp", "vault"} <= caught
    print(f"  {'ok  ' if ok else 'FAIL'} invented terms caught; 'kubernetes' correctly "
          f"NOT flagged (it is in your pool)")
    return ok


# ------------------------------------------------------------- COSTS $ ---- #
@check("expand", "$", "LLM turns an intent into real job titles")
def _expand():
    import llm_tasks as T
    for t in T.expand_query("data engineer roles"):
        print("   ", t)
    return True


@check("grade", "$", "verified scoring: model cites, code checks the quotes")
def _grade():
    import helpers as H, scoring as S, llm_tasks as T, requests
    r = requests.get("https://boards-api.greenhouse.io/v1/boards/figma/jobs",
                     params={"content": "true"}, timeout=30)
    jd = H._strip_html(next(j for j in r.json()["jobs"]
                            if j["title"] == "Data Engineer")["content"])
    edu = "; ".join(f"{e['degree']}, {e['school']}"
                    for e in H.load_profile()["education"])
    reqs = T.extract_requirements(jd)
    print(f"  {len(reqs)} requirements ({sum(r['must_have'] for r in reqs)} required)\n")
    for v in H.load_resume_variants():
        txt = S.resume_text(v) + " Education: " + edu
        cov = S.verify_and_score(txt, T.cite_evidence(reqs, txt))
        print(f"  {v['resume_id']:<16}{cov['score']:>3}/100  "
              f"{len(cov['covered'])}/{len(reqs)} evidenced"
              + (f", {len(cov['fabricated'])} citation(s) rejected"
                 if cov["fabricated"] else ""))
    return True


@check("select", "$", "LLM picks projects from your pool for a JD")
def _select():
    import helpers as H, llm_tasks as T, requests
    r = requests.get("https://boards-api.greenhouse.io/v1/boards/figma/jobs",
                     params={"content": "true"}, timeout=30)
    job = next(j for j in r.json()["jobs"] if j["title"] == "Data Engineer")
    for p in T.select_content(H._strip_html(job["content"]),
                              H.load_experience_pool()):
        print(f"    {p['name'][:44]:<46}{p['why']}")
    return True


@check("profile", note="what the model is told about you")
def _profile():
    import llm_tasks as T
    s = T.profile_summary()
    print(f"  {len(s)} characters\n")
    print("\n".join("  " + l for l in s.splitlines()[:14]))
    print("  ...")
    return True


def main():
    arg = sys.argv[1] if len(sys.argv) > 1 else None

    if not arg:
        print(__doc__)
        print(f"{'CHECK':<14}{'COST':<10}WHAT IT TESTS")
        print("=" * 74)
        for name, (_, cost, note) in CHECKS.items():
            print(f"{name:<14}{cost:<10}{note}")
        return

    names = ([n for n, (_, c, _) in CHECKS.items() if c == "FREE"]
             if arg == "free" else [arg])
    if arg != "free" and arg not in CHECKS:
        print(f"no check called {arg!r}. run with no arguments to list them.")
        return

    failures = []
    for name in names:
        fn, cost, note = CHECKS[name]
        print(f"\n=== {name}  [{cost}] - {note}")
        try:
            if fn() is False:
                failures.append(name)
        except Exception as exc:
            print(f"  ERROR {type(exc).__name__}: {exc}")
            failures.append(name)

    print("\n" + "=" * 74)
    print(f"{len(names) - len(failures)}/{len(names)} passed"
          + (f" - failed: {', '.join(failures)}" if failures else ""))


if __name__ == "__main__":
    main()
